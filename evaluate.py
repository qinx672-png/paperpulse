# -*- coding: utf-8 -*-
"""
步骤 3：LLM 判分评测脚本。

流程：读题库 → 逐题用「被评测模型」作答 → 「判分模型」按参考答案打分 → 汇总报表。

设计要点（对应面试准备）：
- 判分模型与被评测模型解耦：判分固定用强模型（默认 v4-pro），避免自评偏袒
- 防「流畅但错误」放水：判分提示词要求先拆要点、逐点核对，冲突/编造比漏答更严重
- 断点续跑：每题判完即写 JSONL，中断后重跑自动跳过已完成题目
- 三模型对比：--model 一键切换 v4-flash / v4-pro / v3

用法：
    python evaluate.py                          # 全量跑 v4-flash，判分用 v4-pro
    python evaluate.py --model v4-pro           # 换被评测模型
    python evaluate.py --limit 3                # 试水：只跑前 3 题
    python evaluate.py --rows 2-14              # 只跑指定 Excel 行
"""
import argparse
import io
import json
import os
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import pandas as pd

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace", write_through=True)

from config import CHAT_MODELS
from llm import chat
from qa import answer

# ---------- 路径与映射 ----------
TESTBANK = r"C:\Users\17784\research-ai-assistant\科研文献助手人机交互资料\评测题库.xlsx"
RESULT_JSONL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "eval_results.jsonl")
PDF_DIR = r"C:\Users\17784\research-ai-assistant\科研文献助手人机交互资料\题库论文PDF"

# 题库「原文链接」→ 本地 PDF 路径（第 15~20 行链接只写了泛域名，用完整字符串精确匹配）
PDF_MAP = {
    "10.1016/j.scitotenv.2023.169091":
        r"C:\Users\17784\Desktop\YRW和组会\文献精读笔记\YRW-潮汐影响氧化还原敏感元素的释放.pdf",
    "10.1038/s44221-024-00384-6":
        r"C:\Users\17784\Desktop\YRW和组会\文献精读笔记\The hidden influence of terrestrial groundwater on salt marsh function and resilience.pdf",
    "www.nature.com/naturecommunications":
        os.path.join(PDF_DIR, "P1_Huang2017_NatComm_土壤湿度碳矿化.pdf"),
    # 题库改版后 P1 的链接换成了完整 DOI（原来只写泛域名，对不上号，6 题会被跳过）
    "10.1038/s41467-017-01998-z":
        os.path.join(PDF_DIR, "P1_Huang2017_NatComm_土壤湿度碳矿化.pdf"),
    "10.1007/s12647-024-00786-7":
        os.path.join(PDF_DIR, "P2_MAPAN_红树林叶片同位素.pdf"),
    "10.1038/s41467-024-53413-z":
        os.path.join(PDF_DIR, "P3_NatComm_红树林SOC全球.pdf"),
    "10.1111/2041-210X.12183":
        os.path.join(PDF_DIR, "P4_MEE2014_SPOM酸处理.pdf"),
}

# 达标线（用户定稿的评测指标）
PASS_LINE = {"归纳型": 85, "细节型": 75, "图表型": 90}


def find_pdf(link):
    """题库链接 → 本地 PDF 路径。找不到返回 None（如 P4 还没下载）。"""
    link = str(link).strip()
    for key, pdf in PDF_MAP.items():
        if link == key or key in link:
            return pdf if os.path.exists(pdf) else None
    return None


# ---------- 判分提示词 ----------
_JUDGE_SYSTEM = (
    "你是严谨的科研评测判分员，为「PaperPulse（科研文献 AI 助手）」的答题质量打分。\n"
    "判分铁律：\n"
    "1. 【给分要点】就是要点清单（可能用 ①②③、分号、换行分隔），逐点核对即可。\n"
    "   不要自己另拆、合并、增删要点；若【给分要点】为空，再从【完整参考答案】拆解。\n"
    "2. 逐要点核对【助手回答】：\n"
    "   - 答对：要点被正确陈述（数值允许 ±5% 误差；术语中英文都算对）\n"
    "   - 答错：内容与【完整参考答案】冲突，或编造了答案里不存在的事实（比漏答更严重）\n"
    "   - 漏答：完全没有提及\n"
    "3. 警惕「流畅但错误」：表达再通顺，事实与参考答案冲突就是错的，绝不因流畅给分。\n"
    "4. 助手回答「论文未明确提及/检索不到」而参考答案确有该内容 → 该要点记漏答。\n"
    "5. 图表题的特殊评分规则（产品定位：助手不读图，定位+给图表说明才是核心能力）：\n"
    "   - 图内数据要点：要点是图/表里画出来的数据、排序或趋势（如「哪种元素通量最高」\n"
    "     答「V」），而图表说明文字和正文里都没有写出来 → 助手答不出是正确行为。\n"
    "     助手只要正确给出相关图/表编号和页码（或诚实说明无法读图并给出「翻到原图\n"
    "     第 X 页查看」的页码指引），该要点即算【答对】，绝不记为漏答或答错。\n"
    "   - 图表说明文字要点：要点来自 caption 或表格文字（编号、页码、图注描述、\n"
    "     统计说明、表格里的数值）→ 这些是助手材料里有的，仍按逐句核对，\n"
    "     漏复述/漏引用记漏答，复述错误记答错。\n"
    "6. 得分 = 答对要点数 ÷ 要点总数 × 100（四舍五入取整数）；要点总数至少 1 个。\n"
    "只输出 JSON，不要任何多余文字：\n"
    '{"score": 整数, "hit_points": ["..."], "missed_points": ["..."], "wrong_points": ["..."], "reason": "一句话总评"}'
)


def _parse_json(text):
    """从模型输出里抠出 JSON（容忍 ```json 包裹、前后杂字）。"""
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip(), flags=re.M)
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end <= start:
        return None
    try:
        return json.loads(text[start:end + 1])
    except json.JSONDecodeError:
        return None


def judge(question, ref_points, ref_full, model_answer, qtype, judge_model):
    """LLM 判分。解析失败会把上次输出一并附上再问一次，两次都失败返回 None。
    ref_points：题库「参考」列，人工拆好的给分要点；ref_full：题库「答案」列，完整参考答案。
    判分以要点清单为准（分母稳定），完整答案用于判断冲突/编造。"""
    user = (f"【问题】{question}\n"
            f"【给分要点】{ref_points}\n"
            f"【完整参考答案】{ref_full}\n"
            f"【助手回答】{model_answer}")
    prev = ""
    for attempt in range(2):
        extra = f"\n\n上一次你输出的内容无法解析成 JSON，请严格只输出 JSON：{prev}" if prev else ""
        out = chat([{"role": "system", "content": _JUDGE_SYSTEM},
                    {"role": "user", "content": user + extra}],
                   model=judge_model, temperature=0, max_tokens=3000)
        parsed = _parse_json(out)
        if parsed and "score" in parsed:
            return parsed
        prev = out[:400]
    return None


# ---------- 主流程 ----------
def load_done(result_path=RESULT_JSONL):
    """读已完成记录，返回 {(Excel行号, 模型名)} 用于断点续跑。"""
    done = set()
    if os.path.exists(result_path):
        with open(result_path, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    rec = json.loads(line)
                    done.add((rec["row"], rec["model"]))
                except json.JSONDecodeError:
                    continue
    return done


_write_lock = threading.Lock()


def save_record(rec, result_path=RESULT_JSONL):
    with _write_lock:
        with open(result_path, "a", encoding="utf-8") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")


def _do_one(excel_row, row, model, judge_model, result_path, use_rerank=None):
    """单题流水线：作答 → 判分 → 存结果。返回状态字典，供主线程汇总打印。
    use_rerank：rerank 对照实验开关（None=按配置，False=旧基线，True=精排开）。"""
    question = str(row["问题"]).strip()
    # 兼容新旧两版题库列名（2026-09-02）：
    # 新版：答案=要点式参考答案（作给分要点，判分分母），原文参考=原文出处（判冲突/编造）
    # 旧版：答案=完整参考答案（背景），参考=人工拆好的给分要点（评分依据）
    if "原文参考" in row.index:
        ref_points = str(row["答案"]).strip()
        ref_full = str(row["原文参考"]).strip()
    else:
        ref_full = str(row["答案"]).strip()
        ref_points = str(row["参考"]).strip()
    if ref_full == "nan":
        ref_full = ref_points   # 参考列为空时用要点兜底
    qtype = str(row["问题类型"]).strip()

    try:
        t_ans = time.time()
        r = answer(question, row["pdf"], model=model, use_rerank=use_rerank)   # 被评测模型作答
        cost_ans = round(time.time() - t_ans, 1)

        verdict = judge(question, ref_points, ref_full, r["answer"], qtype, judge_model)
    except Exception as e:
        # 网络/API 偶发故障不拖垮整轮：这题不存记录（重跑时自动补），只打警告
        return {"excel_row": excel_row, "ok": False, "verdict": None,
                "route": "异常", "answer": str(e)[:200], "qtype": qtype,
                "error": str(e)[:120]}
    ok = verdict is not None
    if verdict is None:
        verdict = {"score": None, "reason": "判分解析失败"}

    rec = {"row": excel_row, "model": model, "judge_model": judge_model,
           "qtype": qtype, "scene": row["场景层次"] if pd.notna(row["场景层次"]) else "",
           "question": question, "ref": ref_full, "ref_points": ref_points,
           "route": r["route"],
           "answer": r["answer"], "verdict": verdict,
           "cost_ans_s": cost_ans}
    save_record(rec, result_path)
    return {"excel_row": excel_row, "ok": ok, "verdict": verdict, "route": r["route"],
            "answer": r["answer"], "qtype": qtype}


def run(model, judge_model, limit, rows, workers, verbose, testbank=TESTBANK, result_path=RESULT_JSONL,
        use_rerank=None):
    """逐题：作答 → 判分 → 存结果（多线程并发，API 限流很低不会触发）。
    use_rerank：rerank 对照实验开关（None=按配置，False=旧基线，True=精排开）。"""
    df = pd.read_excel(testbank)
    done = load_done(result_path)

    # 先过一遍：筛行、查 PDF、算待跑清单
    jobs = []
    n_skip = 0
    for idx, row in df.iterrows():
        excel_row = idx + 2
        if rows and excel_row not in rows:
            continue
        if (excel_row, model) in done:
            n_skip += 1
            continue
        pdf = find_pdf(row["原文链接"])
        if not pdf:
            print(f"⚠️  第{excel_row}行 PDF 未就绪，跳过（检查 PDF_MAP / P4 是否已下载）")
            continue
        row = row.copy()
        row["pdf"] = pdf
        jobs.append((excel_row, row))

    if limit:
        jobs = jobs[:limit]
    n_fail = 0
    t0 = time.time()

    with ThreadPoolExecutor(max_workers=workers) as ex:
        futures = {ex.submit(_do_one, r, row, model, judge_model, result_path, use_rerank): r for r, row in jobs}
        for i, fut in enumerate(as_completed(futures), 1):
            st = fut.result()
            if "error" in st:
                n_fail += 1
                print(f"[{i}/{len(jobs)}] 第{st['excel_row']}行 ⚠️ 异常跳过（未存记录，重跑自动补）：{st['error']}")
                continue
            if not st["ok"]:
                n_fail += 1
            v = st["verdict"]
            line = f"[{i}/{len(jobs)}] 第{st['excel_row']}行 {st['qtype']} 得分={v.get('score')} | {v.get('reason', '')[:55]}"
            print(line)
            if verbose:
                print(f"   分流={st['route']} | 回答：{st['answer'][:120]}…")

    used = round(time.time() - t0, 1)
    print(f"\n本轮完成：新判 {len(jobs)} 题、续跑跳过 {n_skip} 题、判分失败 {n_fail} 题，用时 {used} 秒")


# ---------- 汇总报表 ----------
def summarize(out_path, judge_model, result_path=RESULT_JSONL):
    """读 JSONL → 按题型 × 模型算均分 → 写 Excel（明细 + 汇总两个 sheet）。"""
    records = []
    with open(result_path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                records.append(json.loads(line))

    if not records:
        print("没有评测记录，先跑 run 再汇总")
        return

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "逐题明细"
    headers = ["行号", "模型", "题型", "场景", "问题", "分流", "得分", "判分理由",
               "命中要点", "漏答要点", "错误要点", "助手回答", "参考答案", "给分要点"]
    ws.append(headers)
    for rec in sorted(records, key=lambda x: (x["row"], x["model"])):
        v = rec.get("verdict") or {}
        ws.append([rec["row"], rec["model"], rec["qtype"], rec.get("scene", ""),
                   rec["question"], rec.get("route", ""), v.get("score"),
                   v.get("reason", ""), "；".join(v.get("hit_points", [])),
                   "；".join(v.get("missed_points", [])), "；".join(v.get("wrong_points", [])),
                   rec["answer"], rec["ref"], rec.get("ref_points", "")])

    # 表头样式 + 列宽
    head_fill = PatternFill("solid", fgColor="D9E2F3")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    widths = [6, 10, 8, 6, 40, 6, 6, 30, 30, 30, 30, 60, 40, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # 汇总 sheet：题型 × 模型均分 + 达标线
    ws2 = wb.create_sheet("题型汇总")
    types = ["归纳型", "细节型", "图表型"]
    models = sorted({r["model"] for r in records})
    ws2.append(["题型"] + models + ["达标线", "是否达标（按最高分模型）"])
    for t in types:
        row_out = [t]
        best = None
        for m in models:
            scores = [r["verdict"]["score"] for r in records
                      if r["qtype"] == t and r["model"] == m
                      and isinstance(r["verdict"].get("score"), (int, float))]
            avg = round(sum(scores) / len(scores), 1) if scores else None
            row_out.append(avg if avg is not None else "—")
            if avg is not None and (best is None or avg > best):
                best = avg
        row_out.append(PASS_LINE[t])
        row_out.append("✅ 达标" if best is not None and best >= PASS_LINE[t]
                       else ("❌ 未达标" if best is not None else "无数据"))
        ws2.append(row_out)

    # 总览行
    all_scores = [r["verdict"]["score"] for r in records
                  if isinstance(r["verdict"].get("score"), (int, float))]
    ws2.append([])
    ws2.append(["全部题目均分（各模型合计）"] + ["" for _ in models] + ["", ""])
    ws2.append(["总体"] + [round(sum(r["verdict"]["score"] for r in records
                                    if r["model"] == m and isinstance(r["verdict"].get("score"), (int, float)))
                                  / max(1, len([r for r in records if r["model"] == m
                                                and isinstance(r["verdict"].get("score"), (int, float))])), 1)
                           if any(r["model"] == m for r in records) else "—"
                           for m in models] + ["", ""])

    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    ws2.column_dimensions["A"].width = 16
    for col in "BCD":
        ws2.column_dimensions[col].width = 12

    wb.save(out_path)
    print(f"✅ 汇总已写：{out_path}")
    print(f"   总记录 {len(records)} 条 | 判分模型 {judge_model}")
    for t in types:
        for m in models:
            scores = [r["verdict"]["score"] for r in records
                      if r["qtype"] == t and r["model"] == m
                      and isinstance(r["verdict"].get("score"), (int, float))]
            if scores:
                print(f"   {t} × {m}：{len(scores)} 题，均分 {round(sum(scores)/len(scores), 1)}（达标线 {PASS_LINE[t]}）")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="PaperPulse LLM 判分评测")
    ap.add_argument("--model", default="v4-flash", choices=list(CHAT_MODELS), help="被评测模型")
    ap.add_argument("--judge", default="v4-pro", choices=list(CHAT_MODELS), help="判分模型（建议固定强模型）")
    ap.add_argument("--limit", type=int, default=0, help="只跑前 N 题（试水用）")
    ap.add_argument("--rows", type=str, default="", help="只跑指定 Excel 行，如 2-14 或 2,5,8")
    ap.add_argument("--workers", type=int, default=3, help="并发数（API 响应慢时提速用）")
    ap.add_argument("--out", default=r"C:\Users\17784\research-ai-assistant\科研文献助手人机交互资料\评测结果.xlsx")
    ap.add_argument("--testbank", default=TESTBANK, help="题库文件路径（修复前后对比时指备份题库）")
    ap.add_argument("--result", default=RESULT_JSONL, help="结果 JSONL 路径（不同轮次分开存）")
    ap.add_argument("--verbose", action="store_true", help="打印回答全文")
    args = ap.parse_args()

    rows = set()
    if args.rows:
        for part in args.rows.split(","):
            part = part.strip()
            if not part:
                continue
            if "-" in part:
                a, b = part.split("-")
                rows.update(range(int(a), int(b) + 1))
            else:
                rows.add(int(part))

    model = CHAT_MODELS[args.model]
    judge_model = CHAT_MODELS[args.judge]
    print(f"被评测模型：{model}")
    print(f"判分模型：{judge_model}")
    if rows:
        print(f"限定行：{sorted(rows)}")
    run(model, judge_model, args.limit, rows, args.workers, args.verbose,
        testbank=args.testbank, result_path=args.result)
    summarize(args.out, judge_model, result_path=args.result)
