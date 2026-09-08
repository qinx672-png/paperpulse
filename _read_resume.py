# -*- coding: utf-8 -*-
"""临时脚本（用完可删）：读取求职文件夹里的 AI 产品经理简历全文。"""
import sys, glob, os
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

folder = r"C:\Users\17784\Desktop\2026-秦肖求职简历合集"
os.chdir(folder)

# 找「AI产品经理-简历」docx（排除作品集/其他方向）
candidates = [f for f in glob.glob("*.docx")
              if "AI" in f and "简历" in f and "作品集" not in f]
print("找到的简历文件：", candidates, "\n" + "=" * 40)

from docx import Document
for name in candidates:
    print(f"\n########## {name} ##########")
    doc = Document(name)
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            print(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            print(" | ".join(cells))
